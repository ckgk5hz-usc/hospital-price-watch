#!/usr/bin/env python3
"""Build data/outcomes.json and data/story_counties.json: what happens when health costs outrun households.

  SHADAC State Health Compare (University of Minnesota), downloaded per indicator as CSV:
    178  adults who could not get medical care when needed due to cost   (CDC BRFSS, annual 2011-2024)
    289  people who had trouble paying off medical bills                  (NHIS, two-year windows 2019-20 .. 2023-24)
    300  prescribed people who changed medicines because of cost          (NHIS, two-year windows)
    222  people in families spending more than 10% of income on health care incl. premiums (CPS ASEC, 2017-2024)
    229  median family medical out-of-pocket spending incl. premiums     (CPS ASEC, 2017-2024)
  Urban Institute, Debt in America 2025 (credit bureau data, August 2025): state and county medical debt in collections.

Exits non-zero on a missing indicator or a changed layout. Never fills a gap. Raw files are cached in tools/data/.cache.
"""
import csv, io, json, re, sys, urllib.request, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CACHE = Path(__file__).resolve().parent / '.cache'
CACHE.mkdir(exist_ok=True)
UA = {'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/120 Safari/537.36'}
fails = []

SHADAC = {
    'care_cost': (178, 'Adults who could not get medical care when needed because of cost', 'BRFSS'),
    'bills': (289, 'People who had trouble paying off medical bills in the past year', 'NHIS'),
    'drugs': (300, 'People prescribed medicine who changed it because of cost', 'NHIS'),
    'burden': (222, 'People in families spending more than 10% of income on health care, premiums included', 'CPS ASEC'),
    'oop': (229, 'Median family out-of-pocket health spending, premiums included', 'CPS ASEC'),
}
URB = 'https://urban-data-catalog.s3.amazonaws.com/drupal-root-live/2025/11/19/Debt%20in%20America%20'


def fetch(url, name, minsize=2000):
    p = CACHE / name
    if not p.exists() or p.stat().st_size < minsize:
        p.write_bytes(urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120).read())
    return p


def num(s):
    s = (s or '').strip()
    return None if s in ('', 'NA', 'N/A') else float(s)


def shadac():
    out = {}
    meta = {}
    for key, (ind, label, survey) in SHADAC.items():
        p = fetch(f'https://statehealthcompare.shadac.org/rawdata.axd?ind={ind}', f'shadac_{ind}.csv')
        rows = list(csv.reader(io.StringIO(p.read_text(encoding='utf-8-sig'))))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == 'Fips')
        assert rows[hi][:6] == ['Fips', 'Location', 'TimeFrame', 'Data Type', 'Data', 'MOE'], f'layout changed for {ind}'
        desc = next((' '.join(r[1:]) for r in rows[:hi] if r and r[0] == 'Description'), '')
        meta[key] = {'ind': ind, 'label': label, 'survey': survey, 'description': desc.replace('&nbsp;', ' ').replace('&rsquo;', "'").strip(), 'title': rows[0][1], 'url': 'https://statehealthcompare.shadac.org/map/%d/%s' % (ind, re.sub(r'-+', '-', re.sub(r'[^a-z0-9]+', '-', rows[0][1].lower())).strip('-'))}
        frames = sorted({r[2] for r in rows[hi + 1:] if len(r) >= 6})
        for r in rows[hi + 1:]:
            if len(r) < 6:
                continue
            loc = {'Dist. of Columbia': 'District of Columbia'}.get(r[1], r[1])
            d = out.setdefault(loc, {}).setdefault(key, {'frames': frames, 'v': [None] * len(frames), 'moe': [None] * len(frames)})
            j = frames.index(r[2])
            d['v'][j] = num(r[4]); d['moe'][j] = num(r[5])
        for need in ('California', 'United States', 'Texas', 'District of Columbia'):
            if need not in out or key not in out[need] or all(v is None for v in out[need][key]['v']):
                fails.append(f'SHADAC {ind}: no data for {need}')
    return out, meta


def urban():
    import openpyxl
    res = {}
    wb = openpyxl.load_workbook(fetch(URB + 'State-Level%20Medical%20Debt.xlsx', 'urban_med_state.xlsx'), data_only=True).active
    rows = list(wb.iter_rows(values_only=True))
    hdr = rows[0]
    idx = {h: i for i, h in enumerate(hdr)}
    cols = {'share': 'Share with medical debt in collections, All', 'share_color': 'Share with medical debt in collections, Comm of color',
            'share_white': 'Share with medical debt in collections, White comm', 'median': 'Median medical debt in collections, All'}
    for k, c in cols.items():
        if c not in idx:
            fails.append('Urban state file: missing column ' + c)
    for r in rows[1:]:
        name = r[idx['State Name']]
        rec = {k: num(str(r[idx[c]])) for k, c in cols.items()}
        # Urban writes 0 for states whose laws restrict credit reporting of medical debt; 0 is not a measurement.
        rec['reported'] = bool(rec['share']) and rec['share'] > 0
        if not rec['reported']:
            rec = {'reported': False}
        res[name] = rec
    nat = list(openpyxl.load_workbook(fetch(URB + 'National-Level%20Medical%20Debt.xlsx', 'urban_med_nat.xlsx'), data_only=True).active.iter_rows(values_only=True))
    ni = {h: i for i, h in enumerate(nat[0])}
    res['United States'] = {'reported': True, 'share': num(str(nat[1][ni[cols['share']]])), 'share_color': num(str(nat[1][ni[cols['share_color']]])),
                            'share_white': num(str(nat[1][ni[cols['share_white']]])), 'median': num(str(nat[1][ni[cols['median']]]))}
    cty = list(openpyxl.load_workbook(fetch(URB + 'County-Level%20Medical%20Debt.xlsx', 'urban_med_cty.xlsx', 50000), data_only=True).active.iter_rows(values_only=True))
    ci = {h: i for i, h in enumerate(cty[0])}
    counties = {}
    for r in cty[1:]:
        sh = num(str(r[ci[cols['share']]])); md = num(str(r[ci[cols['median']]]))
        if sh is None or md is None or not res.get(r[ci['State Name']], {}).get('reported'):
            continue  # an amount is reported only when at least 50 people have medical debt; unreported states carry 0s
        nm = str(r[ci['County Name']]).split(',')[0]
        counties.setdefault(r[ci['State Name']], []).append([nm, round(sh, 4), None if md is None else round(md)])
    if res['United States']['share'] is None:
        fails.append('Urban national share missing')
    if sum(1 for v in res.values() if v.get('reported')) < 40:
        fails.append('Urban: fewer than 40 reporting states')
    return res, counties


def main():
    sh, meta = shadac()
    ub, counties = urban()
    unreported = sorted(k for k, v in ub.items() if not v.get('reported'))
    out = {'generated': datetime.date.today().isoformat(), 'shadac': sh, 'shadac_meta': meta, 'urban': ub, 'urban_unreported': unreported,
           'urban_note': 'Credit bureau data from August 2025. Share of people with a credit record who have medical debt in collections. Excludes states whose laws restrict reporting of medical debt on credit reports.'}
    if fails:
        print('FAILED:\n  ' + '\n  '.join(fails)); sys.exit(1)
    (ROOT / 'data' / 'outcomes.json').write_text(json.dumps(out, separators=(',', ':'), allow_nan=False))
    (ROOT / 'data' / 'story_counties.json').write_text(json.dumps(counties, separators=(',', ':'), allow_nan=False))
    print('wrote data/outcomes.json and data/story_counties.json; unreported states:', unreported)


if __name__ == '__main__':
    main()
